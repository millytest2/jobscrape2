"""Verifies each generated .docx parses cleanly (as an ATS would) and fits one page.
Also emits a .txt of exactly what a parser sees — paste that into ATS 'paste resume' fields."""
import glob
from docx import Document

NS  = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
CPL = 125          # chars per rendered line, Calibri 10pt across a 7.4in column
BUDGET = 46        # rendered lines that fit one Letter page at 0.4in/0.55in margins

ok = True
for f in sorted(glob.glob("resumes/*.docx")):
    doc = Document(f)
    lines, rendered = [], 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        bullet = p._p.find(f'.//{NS}numPr') is not None
        size = max([r.font.size.pt for r in p.runs if r.font.size] or [10])
        cpl = int((CPL - (6 if bullet else 0)) * 10 / size)
        rendered += max(1, -(-len(t) // cpl))
        lines.append(("- " if bullet else "") + t)
    open(f.replace(".docx", ".txt"), "w").write("\n".join(lines))
    fits = rendered <= BUDGET
    ok = ok and fits and not doc.tables
    print(f"{f.split('/')[-1]:48s} lines={rendered:3d}/{BUDGET}  tables={len(doc.tables)}  "
          f"{'OK' if fits else 'OVER BY %d' % (rendered - BUDGET)}")
print("\nALL CHECKS PASSED" if ok else "\nCHECKS FAILED")
raise SystemExit(0 if ok else 1)
